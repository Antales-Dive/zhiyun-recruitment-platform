"""解析器单元测试：TXT 编码、DOCX、文本 PDF 与 OCR 判定。"""
from pathlib import Path

from app.domains.documents.parsers import parse_resume


def make_text_pdf(path: Path) -> None:
    """用 fpdf2 生成带文本层的 PDF fixture。"""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=18)
    pdf.cell(text="Zhang San Python Engineer")
    pdf.ln()
    pdf.cell(text="Email zhangsan@example.com Phone 13800138000")
    pdf.output(str(path))


def test_parses_txt_with_chinese_encoding(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_bytes("姓名：王五\n邮箱 wang@example.com\n电话 13900000000".encode("gb18030"))

    outcome = parse_resume(path, "txt")

    assert outcome.parser_type == "txt"
    assert not outcome.needs_ocr
    text = "\n".join(block.text for block in outcome.blocks)
    assert "王五" in text
    assert "wang@example.com" in text


def test_parses_docx(tmp_path):
    from docx import Document

    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("姓名：赵六")
    document.add_paragraph("邮箱 zhao@example.com")
    document.add_paragraph("电话 13700000000")
    document.add_paragraph("技能：Python、SQL、数据分析与报表开发")
    document.save(str(path))

    outcome = parse_resume(path, "docx")

    assert outcome.parser_type == "docx"
    text = "\n".join(block.text for block in outcome.blocks)
    assert "赵六" in text


def test_parses_text_pdf(tmp_path):
    path = tmp_path / "resume.pdf"
    make_text_pdf(path)

    outcome = parse_resume(path, "pdf")

    assert outcome.parser_type == "pdf_text"
    assert not outcome.needs_ocr
    text = "\n".join(block.text for block in outcome.blocks)
    assert "Zhang San" in text


def test_blank_pdf_requires_ocr(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with open(path, "wb") as handle:
        writer.write(handle)

    outcome = parse_resume(path, "pdf")

    assert outcome.needs_ocr is True
    assert outcome.parser_type == "pdf_text"


def test_image_requires_ocr_without_provider(tmp_path):
    path = tmp_path / "photo.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)

    outcome = parse_resume(path, "png")

    assert outcome.needs_ocr is True


def test_ocr_provider_errors_when_not_configured(tmp_path):
    path = tmp_path / "photo.jpg"
    path.write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 32)

    outcome = parse_resume(path, "jpeg")

    assert outcome.needs_ocr is True
