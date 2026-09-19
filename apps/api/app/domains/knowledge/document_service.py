import re
from dataclasses import dataclass
from pathlib import Path


class OcrRequiredError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedBlock:
    content: str
    section: str | None = None
    page_number: int | None = None


@dataclass(frozen=True)
class DocumentChunk:
    content: str
    section: str | None
    page_number: int | None


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("DOCUMENT_ENCODING_UNSUPPORTED")


def _parse_plain_text(path: Path) -> list[ParsedBlock]:
    text = _decode_text(path.read_bytes())
    blocks: list[ParsedBlock] = []
    section: str | None = None
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(ParsedBlock(content="\n".join(paragraph).strip(), section=section))
            paragraph.clear()

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            flush()
            continue
        heading = re.match(r"^#{1,6}\s+(.+)$", line)
        if heading:
            flush()
            section = heading.group(1).strip()
            continue
        paragraph.append(line)
    flush()
    return blocks


def _parse_docx(path: Path) -> list[ParsedBlock]:
    from docx import Document

    document = Document(str(path))
    blocks: list[ParsedBlock] = []
    section: str | None = None
    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        if not content:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            section = content
        else:
            blocks.append(ParsedBlock(content=content, section=section))
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        content = "\n".join(row for row in rows if row.strip(" |"))
        if content:
            blocks.append(ParsedBlock(content=content, section=section))
    return blocks


def _parse_pdf(path: Path) -> list[ParsedBlock]:
    from pypdf import PdfReader

    blocks = []
    for page_number, page in enumerate(PdfReader(path).pages, start=1):
        content = (page.extract_text() or "").strip()
        if content:
            blocks.append(ParsedBlock(content=content, page_number=page_number))
    if not blocks:
        raise OcrRequiredError("DOCUMENT_REQUIRES_OCR")
    return blocks


def parse_document(path: Path) -> tuple[str, list[ParsedBlock]]:
    suffix = path.suffix.casefold()
    if suffix in {".txt", ".md"}:
        parser_type, blocks = "TEXT", _parse_plain_text(path)
    elif suffix == ".docx":
        parser_type, blocks = "DOCX", _parse_docx(path)
    elif suffix == ".pdf":
        parser_type, blocks = "PDF", _parse_pdf(path)
    elif suffix in {".png", ".jpg", ".jpeg"}:
        raise OcrRequiredError("DOCUMENT_REQUIRES_OCR")
    else:
        raise ValueError("DOCUMENT_TYPE_UNSUPPORTED")
    if not blocks:
        raise ValueError("DOCUMENT_HAS_NO_TEXT")
    return parser_type, blocks


def chunk_blocks(blocks: list[ParsedBlock], max_chars: int = 1200) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    for block in blocks:
        content = block.content.strip()
        while content:
            if len(content) <= max_chars:
                part, content = content, ""
            else:
                boundary = content.rfind("\n", 0, max_chars)
                if boundary < max_chars // 2:
                    boundary = max_chars
                part, content = content[:boundary], content[boundary:]
            chunks.append(
                DocumentChunk(
                    content=part.strip(),
                    section=block.section,
                    page_number=block.page_number,
                )
            )
    return chunks
